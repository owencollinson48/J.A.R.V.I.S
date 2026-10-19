import hashlib
import os
import platform
import shutil
import subprocess
import time
import uuid
from pathlib import Path


TEXT_EXTENSIONS = {".txt", ".md", ".json", ".csv", ".html", ".css", ".js", ".ts", ".py", ".yaml", ".yml", ".toml", ".xml", ".log", ".ini", ".rst"}
SKIP_DIRECTORIES = {".git", "node_modules", ".venv", "__pycache__", "$Recycle.Bin", "Windows", "AppData"}


class Computer:
    def __init__(self, settings):
        self.settings = settings

    def path(self, value):
        path = Path(value).expanduser()
        if not path.is_absolute():
            path = Path(self.settings.values["roots"][0]) / path
        path = path.resolve()
        roots = [Path(r).resolve() for r in self.settings.values["roots"]]
        if not any(path.is_relative_to(root) for root in roots):
            raise ValueError("This path is outside computer access folders. Add its folder in Settings.")
        # App configuration and backups contain private state, not assistant documents.
        if path == self.settings.data_dir or path.is_relative_to(self.settings.data_dir):
            raise ValueError("JARVIS private data is managed through the app, not file tools")
        return path

    def list_files(self, path=None):
        folder = self.path(path or self.settings.values["roots"][0])
        if not folder.is_dir():
            raise ValueError("Choose a folder")
        entries = []
        for item in sorted(folder.iterdir(), key=lambda p: (not p.is_dir(), p.name.lower())):
            try:
                checked = self.path(str(item))
                stat = checked.stat()
                entries.append({"name": item.name, "path": str(item), "directory": checked.is_dir(),
                                "size": stat.st_size, "modified": stat.st_mtime})
            except (OSError, ValueError):
                continue
            if len(entries) >= 200:
                break
        return {"path": str(folder), "parent": str(folder.parent), "entries": entries,
                "limited": len(entries) >= 200}

    def search_files(self, query, directory=None, search_content=False):
        root = self.path(directory or self.settings.values["roots"][0])
        if not query.strip():
            raise ValueError("Enter a search phrase")
        results, scanned = [], 0
        deadline = time.monotonic() + 8
        for base, directories, filenames in os.walk(root, followlinks=False):
            directories[:] = [d for d in directories if d not in SKIP_DIRECTORIES and not d.startswith(".")]
            for name in filenames:
                scanned += 1
                if scanned > 10000 or time.monotonic() > deadline:
                    return {"results": results, "limited": True, "files_scanned": scanned}
                item = Path(base) / name
                try:
                    item = self.path(str(item))
                    snippet = ""
                    match = query.lower() in name.lower()
                    if search_content and item.suffix.lower() in TEXT_EXTENSIONS and item.stat().st_size < 512000:
                        text = item.read_text(encoding="utf-8", errors="replace")
                        index = text.lower().find(query.lower())
                        if index >= 0:
                            match, snippet = True, text[max(0, index - 80):index + 220]
                    if match:
                        results.append({"path": str(item), "name": name, "snippet": snippet})
                except (OSError, ValueError):
                    continue
                if len(results) >= 40:
                    return {"results": results, "limited": True, "files_scanned": scanned}
        return {"results": results, "limited": False, "files_scanned": scanned}

    def read_file(self, path):
        item = self.path(path)
        if item.stat().st_size > 5_000_000:
            raise ValueError("File is larger than the 5 MB document limit")
        if item.suffix.lower() == ".docx":
            from docx import Document
            document = Document(item)
            text = "\n".join(p.text for p in document.paragraphs)
            for table in document.tables:
                text += "\n" + "\n".join(" | ".join(c.text for c in r.cells) for r in table.rows)
        elif item.suffix.lower() in TEXT_EXTENSIONS or not item.suffix:
            data = item.read_bytes()
            try:
                text = data.decode("utf-16" if data.startswith((b'\xff\xfe', b'\xfe\xff')) else "utf-8-sig")
            except UnicodeError as exc:
                raise ValueError("This file is not UTF-8 or UTF-16 text. Save a copy in UTF-8 to preview it.") from exc
        else:
            raise ValueError("Preview supports text files and Word .docx documents")
        return {"path": str(item), "content": text[:24000], "truncated": len(text) > 24000,
                "fingerprint": self.fingerprint(item)}

    @staticmethod
    def fingerprint(path):
        if not path.exists():
            return None
        if not path.is_file():
            raise ValueError("Expected a file")
        if path.stat().st_size > 100_000_000:
            raise ValueError("File is larger than the 100 MB change limit")
        digest = hashlib.sha256()
        with path.open('rb') as handle:
            for chunk in iter(lambda: handle.read(1_048_576), b''): digest.update(chunk)
        return digest.hexdigest()

    def preview(self, name, args):
        preview = {"summary": name.replace("_", " ").title(), "details": args}
        if name in {"write_document", "edit_document"}:
            path = self.path(args["path"])
            if path.suffix.lower() not in TEXT_EXTENSIONS | {".docx", ""}:
                raise ValueError("Writing supports text and .docx documents")
            if len(args.get("content", "").encode("utf-8")) > 1_000_000:
                raise ValueError("Keep a document under 1 MB")
            if not path.parent.is_dir():
                raise ValueError("Create the parent folder first")
            preview["fingerprint"] = self.fingerprint(path)
            preview["summary"] = f"{'Replace' if path.exists() else 'Create'} {path.name}"
            preview["content"] = args.get("content", "")[:6000]
            preview["docx_note"] = ("Writing a .docx replaces its text and formatting. The original is backed up." if path.suffix.lower() == ".docx" and name == 'write_document' else None)
        elif name == "move_file":
            source, destination = self.path(args["source"]), self.path(args["destination"])
            if not source.is_file():
                raise ValueError("Only files can be moved")
            if destination.exists():
                raise ValueError("Destination already exists")
            preview["fingerprint"] = self.fingerprint(source)
        elif name in {"open_file", "create_folder"}:
            self.path(args["path"])
        elif name == "run_command":
            self.path(args.get("directory") or self.settings.values["roots"][0])
            self.validate_command(args["arguments"])
            preview["summary"] = "Run: " + subprocess.list2cmdline(args["arguments"])
        elif name == "open_application":
            preview["summary"] = "Open " + args["name"]
        preview["resolved_paths"] = [str(self.path(args[key])) for key in ("path", "source", "destination", "directory") if args.get(key)]
        return preview

    @staticmethod
    def validate_command(arguments):
        if not isinstance(arguments, list) or not arguments or len(arguments) > 100 or any(not isinstance(a, str) or "\x00" in a for a in arguments):
            raise ValueError("Provide a nonempty command argument list")

    def write_document(self, path, content, fingerprint=None, cancel=None):
        if cancel: cancel.check()
        item = self.path(path)
        if self.fingerprint(item) != fingerprint:
            raise ValueError("File changed after the approval was requested. Read it again before editing.")
        backup = None
        if item.exists():
            folder = self.settings.data_dir / "backups"
            folder.mkdir(exist_ok=True)
            backup = folder / f"{uuid.uuid4().hex}-{item.name}"
            shutil.copy2(item, backup)
        temp = item.with_name(f".{item.name}.{uuid.uuid4().hex}.tmp")
        try:
            if item.suffix.lower() == ".docx":
                from docx import Document
                document = Document()
                for line in content.split("\n"):
                    document.add_paragraph(line)
                document.save(temp)
            else:
                temp.write_text(content, encoding="utf-8")
            if cancel: cancel.check()
            if self.fingerprint(item) != fingerprint: raise ValueError('File changed while the edit was being prepared')
            os.replace(temp, item)
        finally:
            temp.unlink(missing_ok=True)
        return {"written": str(item), "backup": str(backup) if backup else None, 'after_hash':self.fingerprint(item)}

    def edit_document(self, path, replacements=None, font_name=None, font_size=None, line_spacing=None, fingerprint=None, cancel=None):
        """Exact text edits and Word formatting, preserving existing runs/tables/images."""
        if cancel: cancel.check()
        item = self.path(path)
        if not item.is_file() or self.fingerprint(item) != fingerprint: raise ValueError('Read the current file again before editing it')
        replacements = replacements or []
        if not replacements and not any((font_name, font_size, line_spacing)): raise ValueError('Specify text replacements or Word formatting')
        if len(replacements) > 50: raise ValueError('Use at most 50 exact replacements per edit')
        if font_size is not None and not 6 <= font_size <= 96: raise ValueError('Font size must be 6 to 96 points')
        if line_spacing is not None and not .5 <= line_spacing <= 3: raise ValueError('Line spacing must be 0.5 to 3')
        if any(not replacement['find'] for replacement in replacements): raise ValueError('Replacement text to find cannot be empty')
        if item.suffix.lower() != '.docx':
            if any((font_name, font_size, line_spacing)): raise ValueError('Font and paragraph formatting requires a .docx document')
            content = self.read_file(str(item))
            if content['truncated']: raise ValueError('This file is too long for a full-text replacement. Narrow the edit in another editor.')
            text = content['content']
            for replacement in replacements:
                if replacement['find'] not in text: raise ValueError('Exact text was not found; file left unchanged')
                text = text.replace(replacement['find'], replacement['replace'])
            return self.write_document(str(item), text, fingerprint, cancel)
        from docx import Document
        from docx.shared import Pt
        document = Document(item)
        paragraphs = list(document.paragraphs)
        def table_paragraphs(table):
            for row in table.rows:
                for cell in row.cells:
                    yield from cell.paragraphs
                    for inner in cell.tables: yield from table_paragraphs(inner)
        for table in document.tables: paragraphs.extend(table_paragraphs(table))
        for section in document.sections:
            paragraphs.extend(section.header.paragraphs)
            paragraphs.extend(section.footer.paragraphs)
        # Shared table cells/linked headers can expose the same paragraph more than once.
        paragraphs = list({id(paragraph._p): paragraph for paragraph in paragraphs}.values())
        for replacement in replacements:
            found = False
            for paragraph in paragraphs:
                if cancel: cancel.check()
                needle, value = replacement['find'], replacement['replace']
                combined = ''.join(run.text for run in paragraph.runs)
                positions = [match.start() for match in __import__('re').finditer(__import__('re').escape(needle), combined)]
                for start in reversed(positions):
                    end, offset, inserted = start + len(needle), 0, False
                    for run in paragraph.runs:
                        length = len(run.text)
                        if offset < end and offset + length > start:
                            left, right = max(0, start - offset), min(length, end - offset)
                            run.text = run.text[:left] + (value if not inserted else '') + run.text[right:]
                            inserted = True
                        offset += length
                    found = True
            if not found: raise ValueError('Exact text was not found; document left unchanged')
        for paragraph in paragraphs:
            if line_spacing is not None: paragraph.paragraph_format.line_spacing = line_spacing
            for run in paragraph.runs:
                if font_name: run.font.name = font_name
                if font_size is not None: run.font.size = Pt(font_size)
        folder = self.settings.data_dir / 'backups'
        folder.mkdir(exist_ok=True)
        backup = folder / f'{uuid.uuid4().hex}-{item.name}'
        shutil.copy2(item, backup)
        temp = item.with_name(f'.{item.name}.{uuid.uuid4().hex}.tmp')
        try:
            document.save(temp)
            if cancel: cancel.check()
            if self.fingerprint(item) != fingerprint: raise ValueError('Document changed while the edit was being prepared')
            os.replace(temp, item)
        finally: temp.unlink(missing_ok=True)
        return {'written':str(item), 'backup':str(backup), 'after_hash':self.fingerprint(item), 'preserved':'Existing tables, images, styles and unaffected runs retained'}

    def create_folder(self, path):
        item = self.path(path)
        item.mkdir(parents=True, exist_ok=True)
        return {"created": str(item)}

    def move_file(self, source, destination, fingerprint=None, cancel=None):
        if cancel: cancel.check()
        source, destination = self.path(source), self.path(destination)
        if self.fingerprint(source) != fingerprint:
            raise ValueError("Source changed since approval")
        if destination.exists():
            raise ValueError("Destination already exists")
        created = False
        try:
            # Exclusive destination creation also works across drives.
            with destination.open('xb') as output:
                created = True
                with source.open('rb') as input_file:
                    for chunk in iter(lambda: input_file.read(1_048_576), b''):
                        if cancel: cancel.check()
                        output.write(chunk)
            if self.fingerprint(source) != fingerprint or self.fingerprint(destination) != fingerprint:
                raise ValueError('Source changed or copied content did not verify')
            if cancel: cancel.check()
            shutil.copystat(source, destination)
            source.unlink()
        except BaseException:
            if created and source.exists(): destination.unlink(missing_ok=True)
            raise
        return {"moved": str(source), "destination": str(destination), 'after_hash':fingerprint}

    def open_file(self, path):
        item = self.path(path)
        if not item.exists():
            raise ValueError("File does not exist")
        if platform.system() == "Windows":
            os.startfile(str(item))
        else:
            subprocess.Popen(["open" if platform.system() == "Darwin" else "xdg-open", str(item)],
                             stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        return {"requested_open": str(item)}

    def open_application(self, name, arguments=None):
        arguments = arguments or []
        aliases = {"Windows": {"notepad": "notepad.exe", "calculator": "calc.exe", "explorer": "explorer.exe", "paint": "mspaint.exe"},
                   "Darwin": {}, "Linux": {"calculator": "gnome-calculator", "files": "xdg-open"}}
        executable = aliases.get(platform.system(), {}).get(name.lower(), name)
        self.validate_command([executable, *arguments])
        if platform.system() == "Darwin" and not shutil.which(executable) and not Path(executable).is_file():
            command = ["open", "-a", name, "--args", *arguments]
        else:
            command = [executable, *arguments]
        process = subprocess.Popen(command, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                                   cwd=self.settings.values["roots"][0], shell=False)
        return {"launched": name, "pid": process.pid}

    def run_command(self, arguments, directory=None, cancel=None):
        self.validate_command(arguments)
        folder = self.path(directory or self.settings.values["roots"][0])
        # Temporary files avoid unbounded RAM use from a noisy child process.
        import tempfile
        with tempfile.TemporaryFile() as stdout, tempfile.TemporaryFile() as stderr:
            process = subprocess.Popen(arguments, cwd=folder, shell=False, stdout=stdout, stderr=stderr,
                                       start_new_session=platform.system() != "Windows")
            deadline = time.monotonic() + 30
            try:
                while process.poll() is None:
                    if cancel: cancel.check()
                    if time.monotonic() > deadline: raise ValueError('Command exceeded the 30 second limit and was stopped')
                    time.sleep(.1)
            except BaseException:
                import psutil
                try:
                    parent = psutil.Process(process.pid)
                    for child in parent.children(recursive=True):
                        try: child.kill()
                        except psutil.Error: pass
                    parent.kill()
                except psutil.Error: pass
                process.wait(timeout=10)
                raise
            stdout.seek(0)
            stderr.seek(0)
            return {"exit_code": process.returncode,
                    "stdout": stdout.read(16000).decode("utf-8", errors="replace"),
                    "stderr": stderr.read(8000).decode("utf-8", errors="replace")}

    def drives(self):
        if platform.system() == "Windows":
            return [f"{chr(letter)}:\\" for letter in range(65, 91) if Path(f"{chr(letter)}:\\").exists()]
        return [str(Path.home())]
