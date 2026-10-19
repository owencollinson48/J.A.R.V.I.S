import asyncio
import json
import os
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from jarvis.app import create_app
from jarvis.applications import ApplicationIndex, WindowsApplications, app_entry
from jarvis.browser import BrowserControl
from jarvis.computer import Computer
from jarvis.config import Settings
from jarvis.desktop import Images
from jarvis.execution import Cancelled, PermissionPolicy, RunContext
from jarvis.store import Store


class FakeWindows:
    supported = True
    def __init__(self, entries=None):
        self.entries = entries or [app_entry('Outlook (new)','uri','outlookforwindows:','registered protocol',r'C:\Apps\olk.exe'),
                                   app_entry('Calculator','aumid','Microsoft.WindowsCalculator_test!App','Get-StartApps',r'C:\Apps\CalculatorApp.exe')]
        self.visible, self.activations, self.focused, self.failures = [], [], [], set()
    def discover(self, folders): return self.entries, []
    def windows(self): return self.visible[:]
    def focus(self, identity): self.focused.append(identity); return True
    def activate(self, launcher):
        self.activations.append(launcher)
        if launcher['kind'] in self.failures: raise OSError('Activation failed')
        entry = next(e for e in self.entries if launcher in e['launchers'])
        pid = 123 + len(self.activations)
        self.visible.append({'id':100 + len(self.visible),'title':entry['name'],'pid':pid,'executable':entry['executables'][0] if entry['executables'] else '', 'process':Path(entry['executables'][0].replace('\\','/')).name if entry['executables'] else ''})
        return pid


class QueueProvider:
    def __init__(self, settings): self.script = []; self.seen = []
    def headers(self): return {}
    async def check(self): return {'ok':True}
    async def complete(self, messages, tools):
        self.seen.append(messages)
        if self.script:
            item = self.script.pop(0)
            if callable(item): item = await item()
        else: item = {'role':'assistant','content':'Completed the requested task.'}
        yield {'type':'completion','message':item}


def calls(*items):
    return {'role':'assistant','content':None,'tool_calls':[{'id':f'c{i}-{time.time_ns()}','type':'function','function':{'name':name,'arguments':json.dumps(args)}} for i,(name,args) in enumerate(items)]}


def parsed(response):
    assert response.status_code == 200, response.text
    return [json.loads(line[6:]) for line in response.text.splitlines() if line.startswith('data: ')]


@pytest.fixture
def env(tmp_path):
    app = create_app(tmp_path / 'private',provider_factory=QueueProvider)
    root = tmp_path / 'documents'; root.mkdir()
    app.state.settings.update({'roots':[str(root)],'native_notifications':False})
    adapter = FakeWindows(); app.state.tools.applications.adapter = adapter
    with TestClient(app,base_url='http://127.0.0.1') as client:
        client.get('/'); client.headers['X-Jarvis-Client']='dashboard'
        client.root, client.jarvis, client.adapter = root, app, adapter
        yield client


def chat(client, script, text='Complete my requested workflow'):
    client.jarvis.state.assistant.provider.script = script
    cid = client.post('/api/conversations').json()['id']
    return cid, parsed(client.post(f'/api/conversations/{cid}/chat',json={'content':text}))


def test_one_request_plans_edits_saves_and_launches_without_repeated_approval(env):
    path = env.root / 'assignment.txt'; path.write_text('Engineering draft')
    cid, events = chat(env,[calls(('plan_workflow',{'outcome':'Prepare my assignment and open email','steps':['Read assignment','Fix text and save','Open Outlook']})),
        calls(('read_file',{'path':str(path)})),
        calls(('edit_document',{'path':str(path),'replacements':[{'find':'draft','replace':'assignment'}]})),
        calls(('launch_application',{'name':'Outlook'}))])
    assert not any(event['type']=='approval' for event in events)
    assert path.read_text() == 'Engineering assignment'
    assert env.adapter.activations[0]['kind'] == 'uri'
    activity = env.get('/api/activity').json()
    assert activity['workflows'][0]['status'] == 'completed'
    assert len(activity['workflows'][0]['plan']) == 3
    assert len(activity['actions']) == 4
    assert all(action['status']=='completed' for action in activity['actions'])
    edit = next(a for a in activity['actions'] if a['name']=='edit_document')
    assert edit['result']['undo_action_id']
    assert env.jarvis.state.tools.undo(edit['id'],env.jarvis.state.tools.preflight('undo_file_change',{'action_id':edit['id']}),None)['restored'] == str(path)
    assert path.read_text() == 'Engineering draft'


def test_high_impact_command_waits_but_dependent_routine_step_runs_after_one_confirmation(env):
    path = env.root / 'later.txt'
    cid, events = chat(env,[calls(('run_command',{'arguments':[sys.executable,'-c','print("checked")']}),('write_document',{'path':str(path),'content':'saved after confirmation'}))])
    approval = next(e['approval'] for e in events if e['type']=='approval')
    assert len(env.get(f'/api/conversations/{cid}').json()['approvals']) == 1
    assert not path.exists()
    events = parsed(env.post('/api/approvals/' + approval['id'],json={'approved':True}))
    assert not any(e['type']=='approval' for e in events)
    assert path.read_text() == 'saved after confirmation'


def test_decline_cancels_dependent_steps_in_same_tool_batch(env):
    path = env.root / 'must-not-exist.txt'
    cid, events = chat(env,[calls(('run_command',{'arguments':[sys.executable,'-c','print("not run")']}),('write_document',{'path':str(path),'content':'not authorized after decline'}))])
    approval = next(e['approval'] for e in events if e['type']=='approval')
    parsed(env.post('/api/approvals/' + approval['id'],json={'approved':False}))
    assert not path.exists()
    assert not env.get(f'/api/conversations/{cid}').json()['approvals']


def test_consequential_ui_control_is_identified_from_real_metadata_not_model_supplied_label(env):
    class Desktop:
        def preflight(self,args): return {'target_name':'Send payment','password':False,'signature':'trusted'}
        def act(self,*args): raise AssertionError('Must not click before confirmation')
    env.jarvis.state.tools.desktop = Desktop()
    _, events = chat(env,[calls(('desktop_action',{'action':'click','window_id':1,'element_id':'observed'}))])
    approval = next(e['approval'] for e in events if e['type']=='approval')
    assert approval['preview']['risk'] == 'consequential_ui'


def test_credentials_require_review_and_are_not_read_or_logged_in_preflight(env):
    folder = env.root / '.aws'; folder.mkdir(); path = folder / 'credentials'; path.write_text('synthetic-test-secret')
    _, events = chat(env,[calls(('read_file',{'path':str(path)}))])
    approval = next(e['approval'] for e in events if e['type']=='approval')
    assert approval['preview']['risk'] == 'credentials'
    assert 'synthetic-test-secret' not in json.dumps(events)
    assert 'synthetic-test-secret' not in env.get('/api/activity').text


def test_request_can_be_cancelled_during_a_slow_model_response(env):
    async def slow():
        await asyncio.sleep(30)
        return {'role':'assistant','content':'Too late'}
    env.jarvis.state.assistant.provider.script = [slow]
    cid = env.post('/api/conversations').json()['id']
    with ThreadPoolExecutor(1) as executor:
        future = executor.submit(env.post,f'/api/conversations/{cid}/chat',json={'content':'Open Outlook'})
        deadline = time.monotonic() + 3
        while cid not in env.jarvis.state.assistant.contexts and time.monotonic()<deadline: time.sleep(.01)
        assert env.post('/api/cancel').json()['cancel_requested']
        events = parsed(future.result(timeout=3))
    assert any(e['type']=='cancelled' for e in events)
    assert not env.adapter.activations
    assert env.get('/api/activity').json()['workflows'][0]['status']=='cancelled'


def test_cancelling_pending_confirmation_closes_tool_history_and_preserves_next_chat(env):
    cid, events = chat(env,[calls(('run_command',{'arguments':[sys.executable,'-c','print(1)']}))])
    env.post(f'/api/conversations/{cid}/cancel')
    assert env.get(f'/api/conversations/{cid}').json()['approvals']==[]
    parsed(env.post(f'/api/conversations/{cid}/chat',json={'content':'Never mind; hello'}))
    history = env.jarvis.state.store.model_history(cid)
    results={m['tool_call_id'] for m in history if m['role']=='tool'}
    assert all(call['id'] in results for message in history for call in message.get('tool_calls',[]))


def test_index_supports_store_uri_shortcut_executable_and_remembers_preference(tmp_path):
    store = Store(tmp_path/'index.db'); settings = Settings(tmp_path/'private')
    entries = [app_entry('Classic Outlook','exe',r'C:\Office\OUTLOOK.EXE','App Paths',r'C:\Office\OUTLOOK.EXE'),
               app_entry('Outlook (new)','uri','outlookforwindows:','protocol',r'C:\Apps\olk.exe'),
               app_entry('Calculator','aumid','Microsoft.WindowsCalculator_test!App','Get-StartApps',r'C:\Apps\CalculatorApp.exe'),
               app_entry('Portable Editor','shortcut',r'C:\Start Menu\Editor.lnk','Start Menu')]
    adapter=FakeWindows(entries); index=ApplicationIndex(store,settings,adapter)
    assert index.refresh()['count']==4
    for entry in entries:
        result=index.launch(entry['id'],timeout=.1)
        assert result['verified']
    index.alias('my emails',entries[1]['id'])
    assert index.resolve('my emails')['id']==entries[1]['id']
    before=len(adapter.activations)
    assert index.launch('my emails')['already_running']
    assert len(adapter.activations)==before
    assert json.loads(store.row('SELECT payload FROM applications WHERE id=?',(entries[1]['id'],))['payload'])['preferred']['kind']=='uri'
    store.db.close()


def test_unknown_app_is_not_turned_into_a_shell_command(tmp_path):
    store=Store(tmp_path/'index.db'); settings=Settings(tmp_path/'private'); adapter=FakeWindows()
    index=ApplicationIndex(store,settings,adapter); index.refresh()
    with pytest.raises(ValueError,match='not found'): index.launch('zzzz; del C:\\ /s /q')
    assert adapter.activations==[]
    store.db.close()


def test_activation_without_a_visible_window_is_not_claimed_verified(tmp_path):
    store=Store(tmp_path/'index.db'); settings=Settings(tmp_path/'private'); adapter=FakeWindows()
    adapter.activate=lambda method:123
    index=ApplicationIndex(store,settings,adapter); index.refresh()
    result=index.launch('Outlook',timeout=.01)
    assert result['verified'] is False
    assert not json.loads(store.row('SELECT payload FROM applications WHERE id=?',(result['application_id'],))['payload']).get('preferred')
    store.db.close()


def test_fallback_activation_only_after_concrete_method_failure(tmp_path):
    store=Store(tmp_path/'index.db'); settings=Settings(tmp_path/'private')
    entry=app_entry('Editor','aumid','Editor!App','StartApps')
    entry['launchers'].append({'kind':'shortcut','target':r'C:\Editor.lnk','source':'Start Menu'})
    adapter=FakeWindows([entry]); adapter.failures={'aumid'}
    index=ApplicationIndex(store,settings,adapter); index.refresh()
    assert index.launch('Editor',timeout=.1)['verified']
    assert [m['kind'] for m in adapter.activations]==['aumid','shortcut']
    store.db.close()


def test_steam_manifests_become_registered_protocol_launches(tmp_path):
    library=tmp_path/'steamapps'; library.mkdir()
    (library/'appmanifest_123.acf').write_text('"AppState" { "appid" "123" "name" "Engineering Simulator" }')
    apps=WindowsApplications.steam_entries([tmp_path])
    assert apps[0]['launchers'][0]['target']=='steam://rungameid/123'


def test_word_edit_preserves_run_styles_tables_and_line_spacing(env):
    from docx import Document
    path=env.root/'maths.docx'; document=Document(); paragraph=document.add_paragraph()
    paragraph.add_run('Maths ').bold=True; paragraph.add_run('draft').italic=True
    document.add_table(rows=1,cols=1).cell(0,0).text='Keep this table'; document.save(path)
    computer=env.jarvis.state.tools.computer
    result=computer.edit_document(str(path),[{'find':'Maths draft','replace':'Maths assignment'}],font_name='Arial',font_size=12,line_spacing=1.5,fingerprint=computer.fingerprint(path))
    edited=Document(path)
    assert edited.paragraphs[0].text=='Maths assignment'
    assert edited.paragraphs[0].runs[0].bold
    assert edited.tables[0].cell(0,0).text=='Keep this table'
    assert edited.paragraphs[0].paragraph_format.line_spacing==1.5
    assert Path(result['backup']).is_file()


def test_saved_screenshot_becomes_untrusted_vision_context(env):
    from PIL import Image
    image=env.jarvis.state.tools.images.save(Image.new('RGB',(10,10),'white'))
    cid=env.post('/api/conversations').json()['id']
    store=env.jarvis.state.store
    store.message(cid,{'role':'user','content':'Inspect this window'})
    store.message(cid,{'role':'tool','tool_call_id':'image-fixture','content':json.dumps(image)})
    messages=env.jarvis.state.assistant.context(cid)
    assert messages[-1]['content'][1]['image_url']['url'].startswith('data:image/jpeg;base64,')
    assert 'Untrusted screenshots' in messages[-1]['content'][0]['text']
    assert env.get('/api/screenshots/'+image['image_id']).status_code==200


def test_browser_runs_real_navigate_fill_click_and_screenshot(tmp_path,monkeypatch):
    executable=__import__('shutil').which('chromium')
    if not executable: pytest.skip('System Chromium unavailable; use the managed Edge browser on Windows')
    monkeypatch.setenv('JARVIS_BROWSER_EXECUTABLE',executable)
    class Page(BaseHTTPRequestHandler):
        def do_GET(self):
            self.send_response(200);self.send_header('Content-Type','text/html');self.end_headers()
            self.wfile.write(b'<html><body><label>Search<input id="q"></label><button onclick="document.getElementById(\'result\').innerText=document.getElementById(\'q\').value">Find</button><p id="result">Ready</p></body></html>')
        def log_message(self,*args): pass
    server=ThreadingHTTPServer(('127.0.0.1',0),Page); thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    settings=Settings(tmp_path/'private'); browser=BrowserControl(settings,Images(settings.data_dir))
    try:
        args={'action':'navigate','url':f'http://127.0.0.1:{server.server_port}'}
        result=browser.run('act',args,browser.run('preflight',args))
        field=next(e for e in result['observation']['elements'] if e['role']=='textbox')
        args={'action':'fill','element_id':field['id'],'text':'engineering'}
        result=browser.run('act',args,browser.run('preflight',args))
        button=next(e for e in result['observation']['elements'] if e['name']=='Find')
        args={'action':'click','element_id':button['id']}
        result=browser.run('act',args,browser.run('preflight',args))
        assert 'engineering' in result['observation']['text']
        image=browser.run('act',{'action':'screenshot'})
        assert browser.images.path(image['image_id']).is_file()
        with pytest.raises(ValueError,match='HTTP'): browser.run('act',{'action':'navigate','url':'javascript:alert(1)'})
    finally: browser.close();server.shutdown();server.server_close()


def test_index_merges_bridging_package_protocol_identity(tmp_path):
    store = Store(tmp_path/'index.db'); settings = Settings(tmp_path/'settings')
    protocol = app_entry('Outlook (new)','uri','outlookforwindows:','registry')
    start = app_entry('Outlook (new)','aumid','Microsoft.Outlook!App','StartApps')
    package = app_entry('Microsoft.OutlookForWindows','aumid','Microsoft.Outlook!App','package',r'C:\Outlook\olk.exe')
    package['launchers'].append({'kind':'uri','target':'outlookforwindows:','source':'package'})
    index = ApplicationIndex(store,settings,FakeWindows([protocol,start,package]))
    assert index.refresh()['count'] == 1
    app = index.resolve('Outlook')
    assert {method['kind'] for method in app['launchers']} == {'aumid','uri'}
    assert app['executables'] == [r'C:\Outlook\olk.exe']
    store.db.close()


def test_event_actions_match_saved_log_and_reload_history(env):
    cid,events = chat(env,[calls(('launch_application',{'name':'Outlook'}))])
    action = next(event['action'] for event in events if event['type']=='action')
    completed = next(event for event in events if event['type']=='tool')
    assert action['id'] == completed['action_id']
    assert action['status'] == 'running' and completed['status'] == 'completed'
    history = env.get('/api/conversations/' + cid).json()
    assert history['actions'][0]['id'] == action['id']
    assert history['actions'][0]['status'] == 'completed'


def test_launch_check_and_alias_do_not_require_model(env):
    apps = env.get('/api/applications').json()['applications']
    outlook = next(app for app in apps if app['name']=='Outlook (new)')
    response = env.post(f"/api/applications/{outlook['id']}/test")
    assert response.status_code == 200, response.text
    assert response.json()['verified']
    assert not env.jarvis.state.assistant.provider.seen
    assert env.post('/api/applications/alias',json={'alias':'my mail','application_id':outlook['id']}).json()['name']=='Outlook (new)'
    assert env.get('/api/applications?q=my%20mail').json()['applications'][0]['id'] == outlook['id']


def test_created_file_undo_archives_and_changed_file_undo_refuses(env):
    path = env.root/'created.txt'
    _,events = chat(env,[calls(('write_document',{'path':str(path),'content':'first version'}))])
    action = next(event['action_id'] for event in events if event['type']=='tool')
    runner = env.jarvis.state.tools
    preview = runner.preflight('undo_file_change',{'action_id':action})
    path.write_text('user changed this')
    with pytest.raises(ValueError,match='changed'): runner.undo(action,preview,None)
    assert path.read_text() == 'user changed this'
    path.write_text('first version')
    result = runner.undo(action,preview,None)
    assert not path.exists() and Path(result['archived_created_file']).read_text()=='first version'


def test_mutation_policy_checks_resolved_symlink_target(env):
    secret = env.root/'.env'; secret.write_text('synthetic')
    link = env.root/'ordinary.txt'
    try: link.symlink_to(secret)
    except OSError: pytest.skip('Symlinks unavailable')
    runner = env.jarvis.state.tools
    args = {'path':str(link),'content':'changed'}
    preview = runner.preflight('write_document',args)
    context = RunContext(env.jarvis.state.store,env.post('/api/conversations').json()['id'],'Update my file')
    assert runner.policy.decide('write_document',args,preview,context).risk == 'credentials'


def test_cancellation_stops_only_the_owned_command(env):
    marker = env.root/'child.pid'
    code = "import os,time;open('child.pid','w').write(str(os.getpid()));time.sleep(30)"
    context = RunContext(env.jarvis.state.store,env.post('/api/conversations').json()['id'],'Run the test command')
    with ThreadPoolExecutor(1) as executor:
        future = executor.submit(env.jarvis.state.tools.computer.run_command,[sys.executable,'-c',code],str(env.root),context)
        deadline = time.monotonic()+3
        while not marker.exists() and time.monotonic()<deadline: time.sleep(.01)
        assert marker.exists()
        pid = int(marker.read_text())
        context.event.set()
        with pytest.raises(Cancelled): future.result(timeout=3)
    import psutil
    assert not psutil.pid_exists(pid)
